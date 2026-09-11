"""
Unit tests for Document Parser (PDF, DOCX, TXT, text cleaning, and edge cases).
"""

import io
import pytest
from docx import Document
from pypdf import PdfWriter

from src.document_parser import (
    clean_text,
    extract_pdf_text,
    extract_docx_text,
    extract_txt_text,
    extract_text
)
from src.utils import validate_uploaded_file


def create_sample_docx(text: str) -> io.BytesIO:
    """Helper to generate a real DOCX file in memory."""
    doc = Document()
    doc.add_heading("Sample Academic Paper", level=1)
    doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Score"
    table.rows[1].cells[0].text = "F1-Score"
    table.rows[1].cells[1].text = "0.94"
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def create_sample_pdf(text: str) -> io.BytesIO:
    """Helper to generate a valid PDF file in memory using pypdf."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    # pypdf blank page has no text by default
    bio = io.BytesIO()
    writer.write(bio)
    bio.seek(0)
    return bio


class TestTextCleaner:
    def test_clean_text_removes_null_bytes(self):
        dirty = "Introduction\x00 to Machine Learning\x00."
        cleaned = clean_text(dirty)
        assert "\x00" not in cleaned
        assert cleaned == "Introduction to Machine Learning."

    def test_clean_text_normalizes_whitespace_and_newlines(self):
        dirty = "Paragraph 1.   \n\n\n\n\n   Paragraph 2 with   spaces.  "
        cleaned = clean_text(dirty)
        assert "\n\n\n" not in cleaned
        assert "   " not in cleaned
        assert "Paragraph 1.\n\nParagraph 2 with spaces." == cleaned

    def test_clean_text_handles_empty(self):
        assert clean_text("") == ""
        assert clean_text(None) == ""


class TestTXTParser:
    def test_extract_txt_valid(self):
        sample_text = "Title: Deep Learning Survey\nAbstract: We review transformers."
        raw_bytes = sample_text.encode("utf-8")
        res = extract_txt_text(raw_bytes)
        assert res["error"] == ""
        assert "Deep Learning Survey" in res["text"]
        assert "transformers" in res["text"]

    def test_extract_txt_empty(self):
        res = extract_txt_text(b"")
        assert res["text"] == ""
        assert "empty" in res["error"].lower()


class TestDOCXParser:
    def test_extract_docx_valid(self):
        bio = create_sample_docx("This is a study on neural networks and optimization.")
        res = extract_docx_text(bio)
        assert res["error"] == ""
        assert "neural networks and optimization" in res["text"]
        assert "F1-Score" in res["text"]

    def test_extract_docx_corrupted(self):
        corrupted_bytes = b"This is not a real zip or docx file"
        res = extract_docx_text(corrupted_bytes)
        assert res["text"] == ""
        assert "failed to parse" in res["error"].lower()


class TestPDFParser:
    def test_extract_pdf_scanned_detection(self):
        # A blank PDF page contains 0 text, simulating a scanned or empty PDF
        bio = create_sample_pdf("Non-rendered")
        res = extract_pdf_text(bio)
        assert res["is_scanned"] is True
        assert "OCR is not included in MVP" in res["error"]

    def test_extract_pdf_corrupted(self):
        corrupted = b"%PDF-1.4 corrupted header without end of stream"
        res = extract_pdf_text(corrupted)
        assert res["text"] == ""
        assert "failed to parse" in res["error"].lower()


class TestDispatcherAndValidation:
    def test_extract_text_dispatcher_txt(self):
        bio = io.BytesIO(b"Abstract: Analysis of graph neural networks.")
        result = extract_text(bio, "paper1.txt")
        assert result["filename"] == "paper1.txt"
        assert result["file_type"] == "TXT"
        assert result["word_count"] > 0
        assert result["char_count"] > 0
        assert result["token_count"] > 0

    def test_extract_text_unsupported_format(self):
        bio = io.BytesIO(b"some content")
        result = extract_text(bio, "data.csv")
        assert "Unsupported file format" in result["error"]
        assert result["word_count"] == 0

    def test_validate_uploaded_file_valid(self):
        bio = io.BytesIO(b"Valid content")
        val = validate_uploaded_file(bio, "document.pdf")
        assert val["valid"] is True

    def test_validate_uploaded_file_unsupported_extension(self):
        bio = io.BytesIO(b"Some text")
        val = validate_uploaded_file(bio, "presentation.pptx")
        assert val["valid"] is False
        assert "Unsupported file extension" in val["error"]

    def test_validate_uploaded_file_empty(self):
        bio = io.BytesIO(b"")
        bio.size = 0
        val = validate_uploaded_file(bio, "empty.txt")
        assert val["valid"] is False
        assert "empty" in val["error"].lower()
